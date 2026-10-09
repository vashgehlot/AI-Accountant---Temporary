import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from ledgersync import ledger
from ledgersync.export import MONEY, file_name
from ledgersync.models import ClientFields, RowPatch, Transaction
from ledgersync.store import Store

DAY = dt.date(2026, 10, 6)
CUBE = ClientFields(name="Business Cube Ltd", business_type="limited_company", contact_name="Jenny Clarke")


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "ledgersync.db")


def receipt(**changes):
    fields = dict(direction="out", gross="72.00", vat="12.00", account_code="7502", date=dt.date(2026, 9, 1),
                  description="BT Business, broadband", document_type="receipt", counterparty="BT Business")
    return Transaction(**{**fields, **changes})


def sheets(store, client_id):
    content, name = ledger.export(store, client_id, DAY)
    book = load_workbook(io.BytesIO(content))
    return book["Trial balance"], book["Transactions"], name


def test_the_workbook_has_the_trial_balance_and_the_transactions(store):
    client = store.create_client(CUBE)
    store.add_upload(client.id, "BT-0905.pdf", "pdf", [receipt()])
    tb, txs, name = sheets(store, client.id)
    assert name == "Business Cube Ltd 2026-10-06.xlsx"
    assert tb["A1"].value == "Business Cube Ltd — trial balance — 6 Oct 2026"
    assert [[c.value for c in row] for row in tb.iter_rows(min_row=2)] == [
        ["Code", "Account", "Debit", "Credit"], ["1200", "Bank Current Account", 0, 72],
        ["7502", "Telephone and Internet", 60, 0], ["2201", "Purchase VAT", 12, 0], [None, "Totals", 72, 72]]
    assert (tb["C3"].number_format, tb.freeze_panes) == (MONEY, "A3")
    assert [c.value for c in txs[2]] == ["Date", "Description", "Counterparty", "Document type", "In/Out", "Amount",
                                         "VAT", "Net", "Account", "Account name", "Other side", "Still owed",
                                         "Issues", "Upload", "Edited"]
    assert [c.value for c in txs[3]] == [dt.datetime(2026, 9, 1), "BT Business, broadband", "BT Business", "Receipt",
                                         "Out", 72, 12, 60, "7502", "Telephone and Internet", "1200", None, None,
                                         "BT-0905.pdf", None]
    assert (txs["A3"].number_format, txs["F3"].number_format, txs.freeze_panes) == ("dd/mm/yyyy", MONEY, "A3")


def test_while_rows_need_fixing_the_trial_balance_sheet_says_what(store):
    client = store.create_client(CUBE)
    store.add_upload(client.id, "a.pdf", "pdf", [receipt(account_code="9999")])
    tb, _, _ = sheets(store, client.id)
    assert (tb["A2"].value, tb["A3"].value) == ("The trial balance can't be produced yet:",
                                                "#1: This account isn't in the chart.")


def test_an_edited_row_says_so(store):
    client = store.create_client(CUBE)
    store.add_upload(client.id, "a.pdf", "pdf", [receipt()])
    ledger.change_row(store, client.id, store.rows(client.id)[0].id, RowPatch(account_code="7504"))
    _, txs, _ = sheets(store, client.id)
    assert (txs["I3"].value, txs["O3"].value) == ("7504", "Yes")


def test_the_file_name_drops_characters_file_names_cannot_hold():
    # Review focus: a client name with characters a file name can't hold, or letters outside ASCII.
    assert file_name('A/B: "Trading" Ltd', DAY) == "A B Trading Ltd 2026-10-06.xlsx"
    assert file_name("Café £ Ltd", DAY) == "Café £ Ltd 2026-10-06.xlsx"
    assert file_name("???", DAY) == "Client 2026-10-06.xlsx"


def test_document_text_never_becomes_a_formula_or_breaks_the_export(store):
    # Final review #3: "=HYPERLINK(…)" read from a document became a live formula, and a stray control character
    # failed the whole export.
    client = store.create_client(CUBE.model_copy(update={"name": "=Evil Ltd"}))
    store.add_upload(client.id, "=cmd.pdf", "pdf", [receipt(description='=HYPERLINK("http://example.invalid","x")'),
                                                    receipt(description="Paper\x0cand toner", document_ref="p")])
    tb, txs, _ = sheets(store, client.id)
    assert (txs["B3"].value, txs["B3"].data_type) == ('=HYPERLINK("http://example.invalid","x")', "s")
    assert txs["B4"].value == "Paperand toner"
    assert (tb["A1"].data_type, txs["A1"].data_type, txs["N3"].data_type) == ("s", "s", "s")
